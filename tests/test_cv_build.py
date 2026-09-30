"""Contracts for content preservation and accessible document structure."""

import copy
import json
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

import build_cv as cv


@pytest.fixture
def sample():
	return json.loads(cv.pandoc("-f", "gfm", "-t", "json", text="""# Example CV

## Research Supervision

### Master's Thesis Students

- Thesis student

### 2026: 2 students

- First student, master's
- Second student, undergrad

### 2025: 1 student

- Solo student

## Publications

1. A publication with [a link](https://example.org/paper).
1. A second publication.

https://example.org/a/particularly/long/research/project/url
"""))


def plain(document):
	return " ".join(cv.pandoc("-f", "json", "-t", "plain", text=json.dumps(document)).split())


def test_export_preserves_content_and_source(sample):
	original = copy.deepcopy(sample)
	for target in ("html", "docx"):
		assert plain(cv.export_document(sample, target)) == plain(sample)
	assert sample == original


def test_only_multistudent_years_get_columns(sample):
	html = cv.pandoc("-f", "json", "-t", "html5", text=json.dumps(cv.export_document(sample, "html")))
	root = etree.HTML(html)
	groups = root.xpath('//div[@class="student-list"]')
	assert len(groups) == 1
	assert ["".join(item.itertext()) for item in groups[0].xpath("ul/li")] == [
		"First student, master's", "Second student, undergrad"]
	assert not groups[0].xpath(".//table")
	assert len(root.xpath('//span[@class="cv-url"]')) == 1


def test_docx_has_native_lists_columns_and_embedded_fonts(sample, tmp_path):
	reference, output = tmp_path / "reference.docx", tmp_path / "cv.docx"
	cv.reference_document(reference)
	cv.pandoc("-f", "json", "-t", "docx", "--reference-doc", reference,
		"-o", output, text=json.dumps(cv.export_document(sample, "docx")))
	cv.finish_docx(output)
	doc = Document(output)
	assert len(doc.sections) == 3
	assert [s._sectPr.find(cv.qn("w:cols")).get(cv.qn("w:num"), "1") for s in doc.sections] == ["1", "2", "1"]
	assert not doc.tables
	assert doc.core_properties.language == "en-US"
	assert doc.styles["Body Text"].paragraph_format.alignment == cv.WD_ALIGN_PARAGRAPH.JUSTIFY
	for name in ("Heading 1", "Heading 2", "Heading 3"):
		spacing = doc.styles[name].paragraph_format
		assert spacing.space_before > spacing.space_after
	assert (doc.styles["Heading 1"].paragraph_format.left_indent
		< doc.styles["Heading 2"].paragraph_format.left_indent
		< doc.styles["Heading 3"].paragraph_format.left_indent)
	assert doc.part.numbering_part.element.xpath('.//w:numFmt[@w:val="decimal"]')
	assert any(p._p.xpath("./w:pPr/w:numPr") for p in doc.paragraphs)
	assert "CVColumns" not in " ".join(p.text for p in doc.paragraphs)
	assert "First student, master's" in [p.text for p in doc.paragraphs]
	with ZipFile(output) as archive:
		for i, (_, _, filename) in enumerate(cv.FONT_FACES):
			embedded = archive.read(f"word/fonts/cv_font_{i}.odttf")
			assert cv.obfuscate_font(embedded, cv.font_key(filename)) == (cv.FONTS / filename).read_bytes()
		xml = etree.fromstring(archive.read("word/document.xml"))
		assert xml.xpath('//w:hyperlink', namespaces={"w": cv.W_NS})


def test_section_order_cannot_silently_omit_a_file(tmp_path, monkeypatch):
	monkeypatch.setattr(cv, "ROOT", tmp_path)
	folder = tmp_path / "cv"
	folder.mkdir()
	(folder / "profile.md").write_text("# Name\n")
	(folder / "publications.md").write_text("## Publications\n")
	(folder / "sections.txt").write_text("profile.md\n")
	with pytest.raises(ValueError, match="do not match"):
		cv.assemble_source()
	(folder / "sections.txt").write_text("profile.md\npublications.md\n")
	assert cv.assemble_source() == "# Name\n\n## Publications\n"


def test_renderer_rejects_remote_and_unrelated_local_assets():
	fetcher = cv.LocalAssetFetcher()
	for url in ("https://example.org/font.ttf", (cv.ROOT / "cv/profile.md").as_uri()):
		with pytest.raises(ValueError):
			fetcher.fetch(url)
