#!/usr/bin/env python3
"""Build distribution documents from cv/*.md; never import generated documents."""

import copy
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tempfile
import uuid
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fontTools.ttLib import TTFont
from lxml import etree
from pypdf import PdfReader
from weasyprint import HTML
from weasyprint.urls import URLFetcher
from weasyprint.text.fonts import FontConfiguration


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "output"
FONTS = ROOT / "assets" / "fonts"
MAIN_FONT = "Atkinson Hyperlegible Next"
MONO_FONT = "Atkinson Hyperlegible Mono"
URL_FONT = "IBM Plex Sans Condensed"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
FONT_FACES = [
	(MAIN_FONT, "Regular", "AtkinsonHyperlegibleNext-Regular.ttf"),
	(MAIN_FONT, "Bold", "AtkinsonHyperlegibleNext-Bold.ttf"),
	(MAIN_FONT, "Italic", "AtkinsonHyperlegibleNext-Italic.ttf"),
	(MAIN_FONT, "BoldItalic", "AtkinsonHyperlegibleNext-BoldItalic.ttf"),
	(MONO_FONT, "Regular", "AtkinsonHyperlegibleMono-Regular.ttf"),
	(URL_FONT, "Regular", "IBMPlexSansCondensed-Regular.ttf"),
	(URL_FONT, "Italic", "IBMPlexSansCondensed-Italic.ttf"),
]


def assemble_source():
	"""Read the plain section-order file, with every source included exactly once."""
	root = ROOT / "cv"
	names = (root / "sections.txt").read_text().splitlines()
	if not names or len(names) != len(set(names)):
		raise ValueError("Section order must contain each source exactly once")
	for name in names:
		if Path(name).name != name or not name.endswith(".md"):
			raise ValueError(f"Invalid section filename: {name}")
	if set(names) != {path.name for path in root.glob("*.md")}:
		raise ValueError("Section order and Markdown source files do not match")
	return "\n\n".join((root / name).read_text().strip() for name in names) + "\n"


def pandoc(*arguments, text=None):
	"""ASVS 1.2.5: pass arguments directly, with no shell interpolation."""
	return subprocess.run(
		["pandoc", *map(str, arguments)], input=text, text=True,
		check=True, capture_output=True, cwd=ROOT,
	).stdout


def inline_text(value):
	"""Extract displayed text from Pandoc nodes, excluding link targets."""
	if isinstance(value, list):
		return "".join(inline_text(item) for item in value)
	if not isinstance(value, dict):
		return ""
	name, content = value.get("t"), value.get("c")
	if name == "Str":
		return content
	if name in {"Space", "SoftBreak", "LineBreak"}:
		return " "
	if name in {"Link", "Image", "Span"}:
		return inline_text(content[1])
	if name == "Header":
		return inline_text(content[2])
	if name == "Div":
		return inline_text(content[1])
	if name == "Code":
		return content[1]
	return inline_text(content)


def student_lists(document):
	"""Select yearly supervision lists, leaving thesis and other lists alone."""
	selected = set()
	in_supervision = False
	in_year = False
	for index, block in enumerate(document["blocks"]):
		if block["t"] == "Header":
			level, _, content = block["c"]
			text = inline_text(content)
			if level <= 2:
				in_supervision = level == 2 and text.casefold() == "research supervision"
				in_year = False
			elif level == 3:
				in_year = bool(re.match(r"^\d{4}:", text))
		if in_supervision and in_year and block["t"] == "BulletList":
			if len(block["c"]) > 1:
				selected.add(index)
	return selected


def format_links(node, output_format):
	"""Keep complete link targets; give long visible URLs a narrow typeface."""
	if isinstance(node, list):
		return [format_links(item, output_format) for item in node]
	if not isinstance(node, dict):
		return node
	result = {key: format_links(value, output_format) for key, value in node.items()}
	if result.get("t") == "Link":
		visible = inline_text(result["c"][1])
		if len(visible) > 35 and visible.startswith(("https://", "http://")):
			attrs = ["", ["cv-url"], []]
			if output_format == "docx":
				attrs[2].append(["custom-style", "CV URL"])
			result["c"][1] = [{"t": "Span", "c": [attrs, result["c"][1]]}]
	return result


def export_document(document, output_format):
	"""Add output presentation without changing the source content or order."""
	result = copy.deepcopy(document)
	selected = student_lists(document)
	blocks = []
	for index, original in enumerate(result["blocks"]):
		block = format_links(original, output_format)
		if output_format == "docx" and block["t"] == "Header":
			if block["c"][0] == 1:
				block = {"t": "Div", "c": [
					["", [], [["custom-style", "Title"]]],
					[{"t": "Para", "c": block["c"][2]}],
				]}
			else:
				block["c"][0] -= 1
		if index in selected:
			if output_format == "html":
				block = {"t": "Div", "c": [["", ["student-list"], []], [block]]}
			else:
				for style, body in [("CVColumnsStart", []), (None, [block]), ("CVColumnsEnd", [])]:
					if style:
						blocks.append({"t": "RawBlock", "c": ["openxml",
							f'<w:p><w:pPr><w:pStyle w:val="{style}"/></w:pPr></w:p>']})
					else:
						blocks.extend(body)
				continue
		blocks.append(block)
	result["blocks"] = blocks
	result["meta"]["lang"] = {"t": "MetaString", "c": "en-US"}
	return result


def set_font(style, family=MAIN_FONT, size=11, bold=False):
	style.font.name = family
	style.font.size = Pt(size)
	style.font.bold = bold
	style.font.color.rgb = RGBColor(0, 0, 0)
	rpr = style.element.get_or_add_rPr()
	fonts = rpr.find(qn("w:rFonts"))
	if fonts is not None:
		for attr in list(fonts.attrib):
			if "theme" in attr.lower():
				del fonts.attrib[attr]
		for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
			fonts.set(qn("w:" + attr), family)
	lang = OxmlElement("w:lang")
	lang.set(qn("w:val"), "en-US")
	rpr.append(lang)


def reference_document(path):
	"""Generate Word styles from code so the reference is not an editable source."""
	from docx.enum.style import WD_STYLE_TYPE
	doc = Document()
	section = doc.sections[0]
	section.page_width, section.page_height = Inches(8.5), Inches(11)
	for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
		setattr(section, margin, Inches(0.6))
	section.footer_distance = Inches(0.25)
	for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
		style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
		set_font(style)
		style.paragraph_format.line_spacing = 1.1
		style.paragraph_format.space_after = Pt(3 if name == "Compact" else 5)
		style.paragraph_format.widow_control = True
	for number, size in [(1, 14), (2, 12), (3, 11)]:
		style = doc.styles[f"Heading {number}"]
		set_font(style, size=size, bold=True)
		style.paragraph_format.space_before = Pt(10 if number == 1 else 6)
		style.paragraph_format.space_after = Pt(3)
		style.paragraph_format.keep_with_next = True
		if number == 1:
			border = OxmlElement("w:pBdr")
			bottom = OxmlElement("w:bottom")
			for key, value in {"val": "single", "sz": "4", "space": "3", "color": "444444"}.items():
				bottom.set(qn("w:" + key), value)
			border.append(bottom)
			style.element.get_or_add_pPr().append(border)
	set_font(doc.styles["Title"], size=26)
	for border in doc.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
		border.getparent().remove(border)
	doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
	doc.styles["Title"].paragraph_format.space_after = Pt(6)
	for name, font in [("Verbatim Char", MONO_FONT), ("CV URL", URL_FONT), ("Hyperlink", MAIN_FONT)]:
		style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
		set_font(style, family=font)
		if name in {"Hyperlink", "CV URL"}:
			style.font.color.rgb = RGBColor.from_string("154F83")
			style.font.underline = True
	p = section.footer.paragraphs[0]
	p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
	p.add_run("Neil Voss, curriculum vitae | Page ")
	for instruction, separator in [("PAGE", " of "), ("NUMPAGES", "")]:
		field = OxmlElement("w:fldSimple")
		field.set(qn("w:instr"), instruction)
		p._p.append(field)
		p.add_run(separator)
	set_font(doc.styles["Footer"], size=8)
	doc.save(path)


def finish_docx(path):
	"""Replace private markers with continuous native column sections."""
	doc = Document(path)
	# Pandoc's list definitions must use the same bundled font as the document.
	for level in doc.part.numbering_part.element.xpath(".//w:lvl"):
		fmt = level.find(qn("w:numFmt"))
		if fmt is not None and fmt.get(qn("w:val")) == "bullet":
			level.find(qn("w:lvlText")).set(qn("w:val"), "\u2022")
			for font in level.findall(f"./{{{W_NS}}}rPr/{{{W_NS}}}rFonts"):
				for attribute in list(font.attrib):
					del font.attrib[attribute]
				font.set(qn("w:ascii"), MAIN_FONT)
				font.set(qn("w:hAnsi"), MAIN_FONT)
	for paragraph in doc.paragraphs:
		if paragraph._p.xpath("./w:pPr/w:numPr"):
			paragraph.paragraph_format.space_before = Pt(0)
			paragraph.paragraph_format.space_after = Pt(2)
			paragraph.paragraph_format.keep_together = True
		elif paragraph.style.name in {"Body Text", "First Paragraph"}:
			paragraph.paragraph_format.keep_together = True
	base = copy.deepcopy(doc.element.body.sectPr)
	for paragraph in list(doc.paragraphs):
		ppr = paragraph._p.pPr
		style = ppr.find(qn("w:pStyle")) if ppr is not None else None
		name = style.get(qn("w:val")) if style is not None else None
		if name not in {"CVColumnsStart", "CVColumnsEnd"}:
			continue
		ppr.remove(style)
		props = copy.deepcopy(base)
		kind = props.get_or_add_type()
		kind.set(qn("w:val"), "continuous")
		columns = props.find(qn("w:cols"))
		if columns is None:
			columns = OxmlElement("w:cols")
			props.append(columns)
		columns.set(qn("w:num"), "1" if name == "CVColumnsStart" else "2")
		columns.set(qn("w:space"), "360")
		ppr.append(props)
		paragraph.paragraph_format.space_before = Pt(0)
		paragraph.paragraph_format.space_after = Pt(0)
		paragraph.paragraph_format.line_spacing = Pt(1)
		paragraph.paragraph_format.keep_with_next = name == "CVColumnsStart"
	# The final one-column section must also begin continuously.
	kind = doc.element.body.sectPr.get_or_add_type()
	kind.set(qn("w:val"), "continuous")
	doc.core_properties.title = "Neil R. Voss - Curriculum Vitae"
	doc.core_properties.author = "Neil R. Voss"
	doc.core_properties.language = "en-US"
	doc.save(path)
	embed_fonts(path)


def font_key(filename):
	return uuid.uuid5(uuid.NAMESPACE_URL, hashlib.sha256((FONTS / filename).read_bytes()).hexdigest())


def obfuscate_font(data, key):
	"""ECMA-376 font embedding: XOR the first 32 bytes with reversed GUID bytes."""
	result = bytearray(data)
	mask = key.bytes[::-1]
	for index in range(32):
		result[index] ^= mask[index % 16]
	return bytes(result)


def embed_fonts(path):
	with ZipFile(path) as archive:
		parts = {name: archive.read(name) for name in archive.namelist()}
	table = etree.fromstring(parts["word/fontTable.xml"])
	relpath = "word/_rels/fontTable.xml.rels"
	rels = etree.fromstring(parts[relpath]) if relpath in parts else etree.Element("Relationships", nsmap={None: R_NS})
	for index, (family, face, filename) in enumerate(FONT_FACES):
		font = next((f for f in table if f.get(qn("w:name")) == family), None)
		if font is None:
			font = OxmlElement("w:font")
			font.set(qn("w:name"), family)
			table.append(font)
		key = font_key(filename)
		part = f"fonts/cv_font_{index}.odttf"
		parts["word/" + part] = obfuscate_font((FONTS / filename).read_bytes(), key)
		rid = f"rIdCvFont{index}"
		embed = OxmlElement("w:embed" + face)
		embed.set(qn("r:id"), rid)
		embed.set(qn("w:fontKey"), "{" + str(key).upper() + "}")
		font.append(embed)
		etree.SubElement(rels, f"{{{R_NS}}}Relationship", Id=rid,
			Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font", Target=part)
	parts["word/fontTable.xml"] = etree.tostring(table, xml_declaration=True, encoding="UTF-8", standalone=True)
	parts[relpath] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
	types = etree.fromstring(parts["[Content_Types].xml"])
	ctns = "http://schemas.openxmlformats.org/package/2006/content-types"
	etree.SubElement(types, f"{{{ctns}}}Default", Extension="odttf", ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont")
	parts["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
	settings = etree.fromstring(parts["word/settings.xml"])
	settings.append(OxmlElement("w:embedTrueTypeFonts"))
	parts["word/settings.xml"] = etree.tostring(settings, xml_declaration=True, encoding="UTF-8", standalone=True)
	with ZipFile(path, "w", ZIP_DEFLATED) as archive:
		for name, data in parts.items():
			archive.writestr(name, data)


class LocalAssetFetcher(URLFetcher):
	"""ASVS 5.3.2: the renderer reads only bundled assets, never remote URLs."""

	def fetch(self, url, headers=None):
		from urllib.parse import unquote, urlparse
		parsed = urlparse(url)
		if parsed.scheme != "file" or parsed.netloc not in {"", "localhost"}:
			raise ValueError(f"Rendering requires a local asset: {url}")
		path = Path(unquote(parsed.path)).resolve()
		if not any(path.is_relative_to(ROOT / folder) for folder in ("assets", "styles")):
			raise ValueError(f"Asset outside the bundled directories: {path}")
		return super().fetch(url, headers)


def verify_fonts():
	for entry in json.loads((FONTS / "provenance.json").read_text()):
		path = FONTS / entry["file"]
		if path.parent != FONTS or hashlib.sha256(path.read_bytes()).hexdigest() != entry["sha256"]:
			raise ValueError(f"Font asset checksum mismatch: {path}")
	for _, _, filename in FONT_FACES:
		with TTFont(FONTS / filename) as font:
			if font["OS/2"].fsType & 2:
				raise ValueError(f"Font embedding is restricted: {filename}")


def main():
	verify_fonts()
	OUTPUT.mkdir(exist_ok=True)
	(OUTPUT / "pdf").mkdir(exist_ok=True)
	source = assemble_source()
	document = json.loads(pandoc("-f", "gfm", "-t", "json", text=source))
	with tempfile.TemporaryDirectory(prefix="cv-build-", dir=OUTPUT) as temporary:
		temporary = Path(temporary)
		reference = temporary / "reference.docx"
		reference_document(reference)
		docx_path = temporary / "neil_voss_cv.docx"
		pandoc("-f", "json", "-t", "docx", "--reference-doc", reference,
			"-o", docx_path, text=json.dumps(export_document(document, "docx")))
		finish_docx(docx_path)
		body = pandoc("-f", "json", "-t", "html5", text=json.dumps(export_document(document, "html")))
		html = ('<!doctype html><html lang="en-US"><head><meta charset="utf-8">'
			'<meta name="viewport" content="width=device-width, initial-scale=1">'
			'<title>Neil R. Voss - Curriculum Vitae</title>'
			'<link rel="stylesheet" href="../styles/cv.css"></head><body><main>'
			+ body + '</main></body></html>')
		fonts = FontConfiguration()
		pdf_path = temporary / "neil_voss_cv.pdf"
		HTML(string=html, base_url=OUTPUT.as_uri() + "/",
			url_fetcher=LocalAssetFetcher(allowed_protocols={"file"}, fail_on_errors=True)).write_pdf(
			pdf_path, font_config=fonts, pdf_variant="pdf/ua-1")
		pdf = PdfReader(pdf_path)
		if not pdf.pages or "/StructTreeRoot" not in pdf.trailer["/Root"]:
			raise ValueError("PDF is empty or missing its accessibility structure")
		# Publish only after both formats have built successfully.
		docx_path.replace(OUTPUT / docx_path.name)
		pdf_path.replace(OUTPUT / "pdf" / pdf_path.name)
		(OUTPUT / "neil_voss_cv.html").write_text(html)
		(OUTPUT / "CV.md").write_text(source)
	print(f"Built {len(pdf.pages)} PDF pages and an accessible-structure DOCX in {OUTPUT}")


if __name__ == "__main__":
	main()
